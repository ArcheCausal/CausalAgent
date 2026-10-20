"""多模态文档解析与页面路由契约。

文档职责：验证 CSV、XLSX、PDF、图片和文本解析、Docling 页面恢复、RapidOCR 探测以及局部页面路由。

适用范围：修改 multimodal.parsers、文本切分或页面准备规则时使用；远程视觉、索引发布与检索契约由同目录其他测试覆盖。
"""

from __future__ import annotations

import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from Agent.knowledge_base.multimodal.contracts import (
    IngestionIssue,
    IssueSeverity,
    KnowledgeUnit,
    UnitStatus,
    sha256_bytes,
    stable_id,
)
from Agent.knowledge_base.multimodal.pipeline import MultimodalKnowledgeBaseMaintenance
from Agent.knowledge_base.multimodal.parsers import (
    ParsedDocument,
    ParsedItem,
    _convert_docling_page,
    _rapidocr_probe,
    decide_page_route,
    inspect_source,
    parse_document,
)
from Agent.knowledge_base.rag_runtime import RagRuntimeConfig
from Agent.knowledge_base.multimodal.vision import REQUIRED_MODEL, VisionAnalyzer
from Agent.tool_node.rag_questions import normalize_rag_question_output
from Agent.tool_node.rag_tool_registry import build_rag_tools
from tests.unit.agent.multimodal_fixtures import (
    REPOSITORY_ROOT,
    TEST_LOCAL_EMBEDDING,
    outbound_record as _outbound_record,
    png_bytes as _png_bytes,
    vision_payload as _vision_payload,
)


class MultimodalParsingTests(unittest.TestCase):
    """解析器与页面路由：不访问网络，不切换 active pointer。"""

    def test_csv_parser_keeps_header_and_row_semantics(self) -> None:
        """CSV 行的检索文本必须包含表头、行号和单元格值。"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "data.csv"
            source.write_text("处理,结果\n1,改善\n", encoding="utf-8")
            parsed = parse_document(source, "mineru")
            self.assertEqual(parsed.parser_name, "csv")
            self.assertEqual(parsed.items[0].modality, "table")
            self.assertIn("行：2", parsed.items[0].raw_text)
            self.assertIn("处理：1", parsed.items[0].raw_text)
            self.assertIn("结果：改善", parsed.items[0].raw_text)

    def test_xlsx_parser_keeps_sheet_and_row_semantics(self) -> None:
        """XLSX 行必须包含工作表、行号和对应单元格语义。"""
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "data.xlsx"
            workbook = Workbook(); sheet = workbook.active; sheet.title = "实验"; sheet.append(["处理", "结果"]); sheet.append(["干预", "改善"]); workbook.save(source); workbook.close()
            parsed = parse_document(source, "mineru")
            self.assertEqual(parsed.parser_name, "xlsx")
            self.assertIn("工作表：实验", parsed.items[0].raw_text)
            self.assertIn("行：2", parsed.items[0].raw_text)

    def test_standalone_image_parser_does_not_run_local_rapidocr(self) -> None:
        """生产图片解析只保留资源，不再调用 RapidOCR 生成本地 OCR 文本。"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "sample.png"; source.write_bytes(_png_bytes())
            with patch("Agent.knowledge_base.multimodal.parsers._rapidocr_probe") as rapidocr:
                parsed = parse_document(source, "docling")
            rapidocr.assert_not_called()
            self.assertEqual(parsed.items[0].raw_text, "")

    def test_docling_pdf_parser_returns_page_scoped_text(self) -> None:
        """本地 Docling artifacts 必须能把 PDF 正文转换为一基页码单元。"""
        from reportlab.pdfgen import canvas

        artifacts = REPOSITORY_ROOT / "Agent" / "knowledge_base" / "models" / "docling"
        if not artifacts.is_dir():
            self.skipTest("本地 Docling artifacts 未安装")
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "sample.pdf"
            pdf = canvas.Canvas(str(source)); pdf.drawString(72, 720, "Causal inference sample document"); pdf.save()
            with patch.dict(os.environ, {"MULTIMODAL_DOCLING_ARTIFACTS_DIR": str(artifacts), "MULTIMODAL_DOCLING_BATCH_SIZE": "1"}):
                parsed = parse_document(source, "mineru")
            self.assertEqual(parsed.parser_name, "docling")
            self.assertTrue(parsed.items, parsed.issues)
            self.assertTrue(any("Causal inference" in item.raw_text for item in parsed.items))
            self.assertTrue(any(item.page_number == 1 for item in parsed.items))

    def test_docling_pdf_parser_skips_blank_page_without_failing_source(self) -> None:
        """合法空白页应留下 warning，不能使整份 required PDF 失败。"""
        from reportlab.pdfgen import canvas

        artifacts = REPOSITORY_ROOT / "Agent" / "knowledge_base" / "models" / "docling"
        if not artifacts.is_dir():
            self.skipTest("本地 Docling artifacts 未安装")
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "sample.pdf"
            pdf = canvas.Canvas(str(source)); pdf.drawString(72, 720, "Causal inference"); pdf.showPage(); pdf.showPage(); pdf.save()
            with patch.dict(os.environ, {"MULTIMODAL_DOCLING_ARTIFACTS_DIR": str(artifacts), "MULTIMODAL_DOCLING_BATCH_SIZE": "1"}):
                parsed = parse_document(source, "docling")
            self.assertTrue(parsed.items)
            self.assertFalse(any(issue.severity is IssueSeverity.ERROR for issue in parsed.issues))
            self.assertTrue(any(issue.code == "pdf_page_empty" for issue in parsed.issues))

    def test_docling_page_conversion_retries_once_with_new_converter(self) -> None:
        """converter 状态异常时应重建一次，并返回可供后续页复用的新实例。"""
        failed = MagicMock(); failed.convert.side_effect = RuntimeError("stale converter")
        replacement = MagicMock(); replacement.convert.return_value = "result"
        with patch("Agent.knowledge_base.multimodal.parsers._new_docling_converter", return_value=replacement) as factory:
            result, converter = _convert_docling_page(Path("sample.pdf"), 7, object(), failed)
        self.assertEqual(result, "result")
        self.assertIs(converter, replacement)
        factory.assert_called_once()
        replacement.convert.assert_called_once_with(Path("sample.pdf"), page_range=(7, 7))

    def test_empty_docling_table_uses_pdf_region_fallback(self) -> None:
        """空 TableItem 没有 Docling 页图时仍须用原始 PDF bbox 恢复图片。"""
        from Agent.knowledge_base.multimodal.parsers import _docling_items

        class FakeBBox:
            l, r, t, b = 10, 90, 80, 20
            coord_origin = "BOTTOMLEFT"

        class FakeProvenance:
            page_no = 1
            bbox = FakeBBox()

        class FakeTableItem:
            prov = [FakeProvenance()]

            def export_to_markdown(self, doc: object) -> str:
                return ""

        page = MagicMock(size=MagicMock(width=100, height=100))
        document = MagicMock(pages={1: page})
        document.iterate_items.return_value = [(FakeTableItem(), None)]
        recovery_image = _png_bytes()
        with patch("docling_core.types.doc.TableItem", FakeTableItem), patch("Agent.knowledge_base.multimodal.parsers._docling_page_image_bytes", return_value=None), patch("Agent.knowledge_base.multimodal.parsers._render_pdf_region", return_value=recovery_image) as render:
            items, issues = _docling_items(document, Path("sample.pdf"), target_page=1)

        self.assertEqual(issues, [])
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0].content_kind, "table_recovery")
        self.assertEqual(items[0].asset_bytes, recovery_image)
        render.assert_called_once()
        render_args = render.call_args.args
        self.assertEqual(render_args[:2], (Path("sample.pdf"), 1))
        self.assertAlmostEqual(render_args[2]["x0"], 0.1)
        self.assertAlmostEqual(render_args[2]["y0"], 0.2)
        self.assertAlmostEqual(render_args[2]["x1"], 0.9)
        self.assertAlmostEqual(render_args[2]["y1"], 0.8)

    def test_docling_content_page_without_structure_uses_docling_page_image(self) -> None:
        """封面等整页图片未形成结构化项时仍须进入受审核的图片链。"""
        from Agent.knowledge_base.multimodal.parsers import _docling_page_items

        page_image = _png_bytes()
        with patch("Agent.knowledge_base.multimodal.parsers._docling_items", return_value=([], [])), patch("Agent.knowledge_base.multimodal.parsers._pdf_page_has_content", return_value=True), patch("Agent.knowledge_base.multimodal.parsers._docling_page_image_bytes", return_value=page_image):
            items, issues = _docling_page_items(MagicMock(), Path("cover.pdf"), 2)
        self.assertEqual(issues, ())
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0].content_kind, "page_render")
        self.assertEqual(items[0].page_number, 2)
        self.assertEqual(items[0].asset_bytes, page_image)

    def test_rapidocr_probe_classifies_engine_unavailable_failed_and_no_text(self) -> None:
        """RapidOCR 必须区分依赖缺失、推理失败和无文字三种状态，不能统一返回空串。"""
        fp_available = {"engine": "rapidocr-onnxruntime", "engine_version": "1.4.4", "model_fingerprint": "1.4.4", "status": "available"}
        with patch("Agent.knowledge_base.multimodal.parsers.ocr_fingerprint", return_value={**fp_available, "status": "unavailable"}):
            self.assertEqual(_rapidocr_probe(b"image").status, "engine_unavailable")
        broken_engine = MagicMock(); broken_engine.side_effect = RuntimeError("inference failed")
        with patch("Agent.knowledge_base.multimodal.parsers.ocr_fingerprint", return_value=fp_available), patch("Agent.knowledge_base.multimodal.parsers._RAPID_OCR_ENGINE", broken_engine):
            probe = _rapidocr_probe(b"image")
            self.assertEqual(probe.status, "failed")
            self.assertEqual(probe.failure_type, "RuntimeError")
        empty_result = MagicMock(); empty_result.txts = []; empty_result.scores = []
        empty_engine = MagicMock(); empty_engine.return_value = empty_result
        with patch("Agent.knowledge_base.multimodal.parsers.ocr_fingerprint", return_value=fp_available), patch("Agent.knowledge_base.multimodal.parsers._RAPID_OCR_ENGINE", empty_engine):
            self.assertEqual(_rapidocr_probe(b"image").status, "no_text")

    def test_rapidocr_initialization_failure_is_engine_unavailable(self) -> None:
        """模型或引擎初始化失败必须阻断，不能伪装成单图推理失败。"""
        fp_available = {"engine": "rapidocr-onnxruntime", "engine_version": "1.4.4", "model_fingerprint": "a" * 64, "status": "available"}
        with patch("Agent.knowledge_base.multimodal.parsers.ocr_fingerprint", return_value=fp_available), patch("Agent.knowledge_base.multimodal.parsers._RAPID_OCR_ENGINE", None), patch("rapidocr.RapidOCR", side_effect=FileNotFoundError("model missing")):
            probe = _rapidocr_probe(b"image")
        self.assertEqual(probe.status, "engine_unavailable")
        self.assertEqual(probe.failure_type, "FileNotFoundError")

    def test_ocr_engine_unavailable_produces_blocking_issue(self) -> None:
        """依赖缺失必须生成 ERROR 级 issue，使正式来源在门禁处被阻断。"""
        from Agent.knowledge_base.multimodal.parsers import _ocr_issues, OcrResult
        probe = OcrResult("", "rapidocr-onnxruntime", "1.4.4", "1.4.4", 0, 0.0, 0, "engine_unavailable")
        issues = _ocr_issues(probe, Path("source.pdf"))
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].code, "ocr_engine_unavailable")
        self.assertEqual(issues[0].severity, IssueSeverity.ERROR)
        self.assertTrue(issues[0].blocking)

    def test_split_text_units_prefers_headings_and_keeps_target_size(self) -> None:
        """TXT/Markdown 切分必须优先标题/段落边界，目标约 800、最大 1200、重叠约 120。"""
        from Agent.knowledge_base.multimodal.parsers import TEXT_SPLIT, split_text_units
        self.assertEqual(TEXT_SPLIT["target_chars"], 800)
        self.assertEqual(TEXT_SPLIT["max_chars"], 1200)
        self.assertEqual(TEXT_SPLIT["overlap_chars"], 120)
        short = "短文不切分"
        self.assertEqual(split_text_units(short), [short])
        paragraphs = "\n\n".join(f"段落{index}：" + ("因果推理。" * 90) for index in range(1, 6))
        chunks = split_text_units(paragraphs)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk) <= TEXT_SPLIT["max_chars"] for chunk in chunks))
        self.assertTrue(any(len(chunk) >= TEXT_SPLIT["target_chars"] * 0.6 for chunk in chunks))
        # 每节约 900 字，超过 target；标题边界应阻止并入上一节
        heading_text = "# 标题一\n" + ("正文甲。" * 225) + "\n\n## 标题二\n" + ("正文乙。" * 225)
        headed = split_text_units(heading_text)
        self.assertGreaterEqual(len(headed), 2)
        self.assertTrue(any(chunk.lstrip().startswith("# 标题一") for chunk in headed))
        self.assertTrue(any(chunk.lstrip().startswith("## 标题二") or "标题二" in chunk[:40] for chunk in headed))
        hard = "字" * 2500
        hard_chunks = split_text_units(hard)
        self.assertGreaterEqual(len(hard_chunks), 2)
        self.assertTrue(all(len(chunk) <= TEXT_SPLIT["max_chars"] for chunk in hard_chunks))
        if len(hard_chunks) >= 2:
            overlap = min(TEXT_SPLIT["overlap_chars"], len(hard_chunks[0]), len(hard_chunks[1]))
            self.assertTrue(hard_chunks[0][-overlap:] == hard_chunks[1][:overlap] or hard_chunks[0][-40:] in hard_chunks[1])

    def test_parse_text_emits_multiple_paragraph_units(self) -> None:
        """长 TXT/Markdown 文件必须被切成多个 text/paragraph 单元，而不是单个超长单元。"""
        from Agent.knowledge_base.multimodal.parsers import TEXT_SPLIT
        with tempfile.TemporaryDirectory() as directory:
            body = "\n\n".join(f"## 段{index}\n" + ("因果。" * 100) for index in range(1, 5))
            source = Path(directory) / "notes.md"
            source.write_text(body, encoding="utf-8")
            parsed = parse_document(source, "docling")
            self.assertEqual(parsed.parser_name, "text")
            self.assertGreater(len(parsed.items), 1)
            self.assertTrue(all(item.modality == "text" and item.content_kind == "paragraph" for item in parsed.items))
            self.assertTrue(all(len(item.raw_text) <= TEXT_SPLIT["max_chars"] for item in parsed.items))
            joined = "\n".join(item.raw_text for item in parsed.items)
            self.assertIn("因果", joined)

    def test_page_routes_are_deterministic_and_mutually_exclusive(self) -> None:
        """空白、对象图片和整页回退必须产生唯一且稳定的本地路由。"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "page.png"
            source.write_bytes(_png_bytes())
            local = ParsedDocument("builtin", "1", (ParsedItem("text", "paragraph", raw_text="正文", page_number=1),))
            picture = ParsedDocument("builtin", "1", (ParsedItem("image", "image", asset_bytes=_png_bytes(), asset_name="picture.png", page_number=1),))
            fallback = ParsedDocument("docling", "2", (ParsedItem("page", "page_render", asset_bytes=_png_bytes(), asset_name="page.png", page_number=1),))
            blank = ParsedDocument("builtin", "1")
            self.assertEqual(decide_page_route(source, 1, local).route, "local_objects")
            self.assertEqual(decide_page_route(source, 1, picture).route, "remote_pictures")
            self.assertEqual(decide_page_route(source, 1, fallback).route, "remote_page_fallback")
            self.assertEqual(decide_page_route(source, 1, blank).route, "blank_page")
            self.assertEqual(MultimodalKnowledgeBaseMaintenance._routed_page_items((picture.items[0], fallback.items[0]), "remote_page_fallback"), (fallback.items[0],))

    def test_page_preparation_keeps_unpaired_short_text(self) -> None:
        """质量门后的单页短正文不得因标题合并策略而静默丢失。"""
        service = MultimodalKnowledgeBaseMaintenance()
        items, filtered = service._prepare_page_items((ParsedItem("text", "paragraph", raw_text="短正文", page_number=1),))
        self.assertEqual(filtered, 1)
        self.assertEqual([item.raw_text for item in items], ["短正文"])


if __name__ == "__main__":
    unittest.main()
